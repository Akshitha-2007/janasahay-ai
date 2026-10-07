"""JANASAHAY AI - RAG engine: ingestion, retrieval, grounded generation."""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import uuid
from datetime import datetime
from pathlib import Path

import streamlit as st

BASE = Path(__file__).resolve().parent.parent
DATA = BASE / "data"
UPLOADS, CHROMA, SAMPLES = DATA / "uploads", DATA / "chroma", DATA / "sample_docs"
REG, LOG = DATA / "registry.json", DATA / "queries.json"
for _p in (UPLOADS, CHROMA):
    _p.mkdir(parents=True, exist_ok=True)

EMBED_MODEL = "all-MiniLM-L6-v2"
DEFAULTS = {"state": "All India", "occupation": "Any", "category": "Any"}

SAMPLE_META = {
    "pm_kisan.txt": dict(scheme="PM-KISAN", occupation="Farmer", scheme_type="Agriculture"),
    "mudra_yojana.txt": dict(scheme="PM Mudra Yojana", occupation="Entrepreneur", scheme_type="Business & Loans"),
    "stand_up_india.txt": dict(scheme="Stand-Up India", occupation="Entrepreneur", category="Women / SC / ST", scheme_type="Business & Loans"),
    "national_scholarship_portal.txt": dict(scheme="National Scholarship Portal (Post Matric)", occupation="Student", scheme_type="Education"),
    "pudhumai_penn.txt": dict(scheme="Pudhumai Penn Scheme", state="Tamil Nadu", occupation="Student", category="Women", scheme_type="Education"),
}


class RagError(Exception):
    """User-friendly error shown in the UI."""


def get_secret(name: str) -> str:
    v = os.environ.get(name)
    if v:
        return v
    try:
        return str(st.secrets.get(name, "") or "")
    except Exception:
        return ""


# ---------------------------------------------------------------- embeddings
def _hash_vec(text: str, dim: int = 384) -> list[float]:  # used only when JANASAHAY_TEST_EMBED is set
    v = [0.0] * dim
    for w in re.findall(r"\w+", text.lower()):
        v[int(hashlib.md5(w.encode()).hexdigest(), 16) % dim] += 1
    n = sum(x * x for x in v) ** 0.5 or 1
    return [x / n for x in v]


@st.cache_resource(show_spinner="Loading embedding model (first run only)...")
def get_embedder():
    if os.environ.get("JANASAHAY_TEST_EMBED"):
        return "hash"
    try:
        from sentence_transformers import SentenceTransformer
        return SentenceTransformer(EMBED_MODEL)
    except Exception:
        return None


def embed(texts: list[str]) -> list[list[float]]:
    m = get_embedder()
    if m is None:
        raise RagError("The embedding model (all-MiniLM-L6-v2) could not be loaded. Check the internet connection / dependencies and reload.")
    if m == "hash":
        return [_hash_vec(t) for t in texts]
    return m.encode(texts, normalize_embeddings=True, show_progress_bar=False).tolist()


@st.cache_resource
def get_collection():
    try:
        import chromadb
        client = chromadb.PersistentClient(path=str(CHROMA))
        return client.get_or_create_collection("schemes", metadata={"hnsw:space": "cosine"})
    except Exception as e:
        raise RagError(f"Vector database (ChromaDB) is unavailable: {e}")


# ---------------------------------------------------------------- registry
def _load(path: Path, default):
    try:
        return json.loads(path.read_text("utf-8"))
    except Exception:
        return default


def registry() -> dict:
    return _load(REG, {})


def _save_registry(r: dict):
    REG.write_text(json.dumps(r, indent=1), "utf-8")


def schemes() -> list[str]:
    return sorted({d["scheme"] for d in registry().values()})


def options(field: str) -> list[str]:
    return sorted({d.get(field, "") for d in registry().values() if d.get(field)})


# ---------------------------------------------------------------- extraction
def extract_text(name: str, data: bytes) -> tuple[str, bool]:
    ext = Path(name).suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
            if len(text.strip()) >= 50:
                return text, False
            try:  # scanned PDF -> OCR
                import pytesseract
                from pdf2image import convert_from_bytes
                imgs = convert_from_bytes(data, dpi=200, first_page=1, last_page=30)
                return "\n".join(pytesseract.image_to_string(i) for i in imgs), True
            except Exception:
                raise RagError(f"'{name}' has no extractable text and OCR failed (is Tesseract/Poppler installed?).")
        if ext == ".docx":
            from docx import Document
            d = Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                parts += [" | ".join(c.text for c in r.cells) for r in t.rows]
            return "\n".join(parts), False
        if ext == ".txt":
            return data.decode("utf-8", errors="ignore"), False
    except RagError:
        raise
    except Exception as e:
        raise RagError(f"'{name}' could not be read - the file may be corrupt or invalid ({e}).")
    raise RagError(f"Unsupported file type '{ext}'. Upload PDF, TXT or DOCX.")


def clean_text(t: str) -> str:
    t = t.replace("\x00", " ")
    t = re.sub(r"[ \t]+", " ", t)
    return re.sub(r"\n{3,}", "\n\n", t).strip()


def chunk_text(text: str, size: int = 180, overlap: int = 40) -> list[str]:
    words, out = text.split(), []
    for i in range(0, len(words), size - overlap):
        c = " ".join(words[i:i + size])
        if len(c) > 40:
            out.append(c)
        if i + size >= len(words):
            break
    return out


# ---------------------------------------------------------------- ingestion
def ingest(name: str, data: bytes, meta: dict | None = None, doc_id: str | None = None, progress=None) -> dict:
    step = progress or (lambda f, m: None)
    doc_id = doc_id or uuid.uuid4().hex[:10]
    step(0.10, "Extracting text...")
    text, ocr = extract_text(name, data)
    step(0.30, "Cleaning text...")
    text = clean_text(text)
    if len(text) < 50:
        raise RagError(f"'{name}' contains too little readable text to index.")
    step(0.45, "Chunking...")
    chunks = chunk_text(text)
    step(0.60, f"Embedding {len(chunks)} chunks...")
    vecs = embed(chunks)
    first = next((l.strip() for l in text.splitlines() if l.strip()), Path(name).stem)
    m = {"scheme": first if len(first) < 90 else Path(name).stem, "state": "All India",
         "occupation": "Any", "category": "Any", "scheme_type": "General"}
    m.update({k: v for k, v in (meta or {}).items() if v})
    step(0.85, "Storing in ChromaDB...")
    col = get_collection()
    col.add(ids=[f"{doc_id}-{i}" for i in range(len(chunks))], documents=chunks, embeddings=vecs,
            metadatas=[{**m, "doc_id": doc_id, "file": name, "chunk": i} for i in range(len(chunks))])
    path = UPLOADS / f"{doc_id}_{Path(name).name}"
    path.write_bytes(data)
    reg = registry()
    reg[doc_id] = {**m, "id": doc_id, "file": name, "type": Path(name).suffix.upper().strip("."),
                   "uploaded": datetime.now().strftime("%Y-%m-%d %H:%M"), "chunks": len(chunks),
                   "status": "Indexed", "ocr": ocr, "path": str(path)}
    _save_registry(reg)
    step(1.0, "Document indexed successfully.")
    return reg[doc_id]


def delete_doc(doc_id: str):
    try:
        get_collection().delete(where={"doc_id": doc_id})
    except Exception:
        pass
    reg = registry()
    d = reg.pop(doc_id, None)
    if d:
        Path(d.get("path", "")).unlink(missing_ok=True)
    _save_registry(reg)


def reindex_doc(doc_id: str):
    d = registry().get(doc_id)
    if not d or not Path(d["path"]).exists():
        raise RagError("Original file not found on disk; please upload it again.")
    data = Path(d["path"]).read_bytes()
    meta = {k: d[k] for k in ("scheme", "state", "occupation", "category", "scheme_type")}
    get_collection().delete(where={"doc_id": doc_id})
    Path(d["path"]).unlink(missing_ok=True)
    return ingest(d["file"], data, meta, doc_id)


def load_samples() -> int:
    n = 0
    have = {d["file"] for d in registry().values()}
    for f in sorted(SAMPLES.glob("*.txt")):
        if f.name not in have:
            ingest(f.name, f.read_bytes(), SAMPLE_META.get(f.name))
            n += 1
    return n


def doc_chunks(doc_id: str, limit: int = 50) -> list[str]:
    r = get_collection().get(where={"doc_id": doc_id}, limit=limit)
    pairs = sorted(zip(r["metadatas"], r["documents"]), key=lambda x: x[0]["chunk"])
    return [p[1] for p in pairs]


# ---------------------------------------------------------------- retrieval
def _where(f: dict | None):
    conds = []
    for k, v in (f or {}).items():
        if not v or v in ("Any", "All"):
            continue
        conds.append({k: {"$in": [v, DEFAULTS[k]]}} if k in DEFAULTS else {k: v})
    return None if not conds else conds[0] if len(conds) == 1 else {"$and": conds}


def retrieve(query: str, k: int = 5, filters: dict | None = None) -> list[dict]:
    col = get_collection()
    n = col.count()
    if n == 0:
        return []
    qv = embed([query])
    try:
        res = col.query(query_embeddings=qv, n_results=min(k, n), where=_where(filters))
    except Exception:
        return []
    return [{"text": d, "meta": m, "score": round(max(0.0, 1 - dist), 3)}
            for d, m, dist in zip(res["documents"][0], res["metadatas"][0], res["distances"][0])]


def log_query(q: str, hits: list[dict], kind: str):
    log = _load(LOG, [])
    avg = round(sum(h["score"] for h in hits) / len(hits), 3) if hits else 0.0
    log.append({"ts": datetime.now().isoformat(timespec="seconds"), "q": q[:200], "avg": avg, "kind": kind})
    try:
        LOG.write_text(json.dumps(log[-5000:]), "utf-8")
    except Exception:
        pass


def query_log() -> list[dict]:
    return _load(LOG, [])


def clear_log():
    LOG.unlink(missing_ok=True)


# ---------------------------------------------------------------- generation
RULE = ("Answer ONLY from the retrieved context. If the information is not available in the retrieved context, "
        "clearly state that sufficient information was not found. Never invent scheme names, amounts, dates or eligibility rules.")


def llm(prompt: str, json_mode: bool = False) -> str:
    key = get_secret("GEMINI_API_KEY")
    if not key:
        raise RagError("GEMINI_API_KEY is not configured. Add it to Streamlit secrets or your environment to enable AI answers.")
    try:
        from google import genai
        from google.genai import types
        cfg = types.GenerateContentConfig(temperature=0.1, response_mime_type="application/json" if json_mode else None)
        r = genai.Client(api_key=key).models.generate_content(
            model=get_secret("GEMINI_MODEL") or "gemini-2.5-flash", contents=prompt, config=cfg)
        return r.text or ""
    except Exception as e:
        raise RagError(f"Gemini request failed: {str(e)[:200]}")


def context_block(hits: list[dict]) -> str:
    return "\n\n".join(f"[{i}] Scheme: {h['meta']['scheme']} | File: {h['meta']['file']}\n{h['text']}"
                       for i, h in enumerate(hits, 1))


def parse_json(t: str):
    t = re.sub(r"^```(?:json)?|```$", "", t.strip(), flags=re.M).strip()
    try:
        return json.loads(t)
    except Exception:
        m = re.search(r"[\[{].*[\]}]", t, re.S)
        try:
            return json.loads(m.group(0)) if m else None
        except Exception:
            return None


def answer(question: str, filters: dict | None = None, k: int = 5) -> dict:
    hits = retrieve(question, k, filters)
    log_query(question, hits, "chat")
    if not hits:
        return {"answer": "No relevant documents were found. Please ask an admin to upload / index government scheme documents first.", "hits": []}
    prompt = (f"You are JANASAHAY AI, an assistant for Indian government welfare schemes.\n{RULE}\n"
              "Cite sources like [1], [2]. Use short, clear markdown.\n\n"
              f"CONTEXT:\n{context_block(hits)}\n\nQUESTION: {question}\n\nANSWER:")
    try:
        return {"answer": llm(prompt), "hits": hits}
    except RagError as e:
        return {"answer": f"⚠️ {e}\n\nShowing the most relevant retrieved passages below instead.", "hits": hits}


def eligibility(p: dict) -> dict:
    q = (f"Schemes for {p['gender']} {p['occupation']} aged {p['age']}, education {p['education']}, "
         f"annual income Rs {p['income']}, category {p['category']}, state {p['state']}. {p.get('extra', '')}")
    hits = retrieve(q, 8, {"state": p["state"]})
    log_query(q, hits, "eligibility")
    if not hits:
        return {"schemes": [], "hits": []}
    prompt = (f"{RULE}\nUser profile: {json.dumps(p)}\nUsing ONLY the context, list schemes this person may qualify for. "
              'Return a JSON array of objects with keys: "name","eligibility","benefits","documents","why","source" '
              "(source = file name). Return [] if none clearly match.\n\nCONTEXT:\n" + context_block(hits))
    data = parse_json(llm(prompt, True))
    return {"schemes": data if isinstance(data, list) else [], "hits": hits}


COMPARE_FIELDS = ["Eligibility", "Benefits", "Income Limit", "Documents", "Application Process"]


def compare(names: list[str]) -> dict:
    out = {}
    for n in names:
        hits = retrieve(f"{n} eligibility benefits income limit documents application process", 4, {"scheme": n})
        log_query(f"compare {n}", hits, "compare")
        prompt = (f"{RULE}\nFor the scheme '{n}', return a JSON object with keys "
                  '"Eligibility","Benefits","Income Limit","Documents","Application Process" (strings). '
                  'Use "Not found in documents" when absent.\n\nCONTEXT:\n' + context_block(hits))
        d = parse_json(llm(prompt, True))
        out[n] = d if isinstance(d, dict) else {}
    return out


def checklist(name: str) -> dict:
    hits = retrieve(f"{name} documents required application process", 4, {"scheme": name})
    log_query(f"checklist {name}", hits, "checklist")
    prompt = (f"{RULE}\nFor '{name}' return JSON " '{"documents":[...],"steps":[...]} ' "using only the context.\n\nCONTEXT:\n" + context_block(hits))
    d = parse_json(llm(prompt, True))
    return d if isinstance(d, dict) else {"documents": [], "steps": []}


# ---------------------------------------------------------------- evaluation / stats
def default_cases() -> list[tuple[str, str]]:
    cs = []
    for s in schemes():
        cs += [(f"Who is eligible for {s}?", s), (f"What are the benefits and required documents of {s}?", s)]
    return cs


def evaluate(cases: list[tuple[str, str]], k: int = 5) -> list[dict]:
    rows = []
    for q, target in cases:
        hits = retrieve(q, k)
        rel = sum(1 for h in hits if h["meta"]["scheme"].lower() == target.lower())
        p = rel / len(hits) if hits else 0.0
        r = 1.0 if rel else 0.0
        f1 = 2 * p * r / (p + r) if p + r else 0.0
        rows.append({"Question": q, "Expected scheme": target, "Relevant@K": rel,
                     "Precision@K": round(p, 3), "Recall@K": r, "F1": round(f1, 3)})
    return rows


def stats() -> dict:
    reg, log = registry(), query_log()
    try:
        chunks = get_collection().count()
    except RagError:
        chunks = 0
    last = max(reg.values(), key=lambda d: d["uploaded"]) if reg else None
    scored = [q["avg"] for q in log if q["avg"]]
    return {"docs": len(reg), "chunks": chunks, "schemes": len({d["scheme"] for d in reg.values()}),
            "queries": len(log), "avg_score": round(sum(scored) / len(scored), 3) if scored else 0.0, "last": last}
